import openpyxl
from collections import Counter

path = "/Users/michelleyip/Documents/PM-OS/context-library/research/OTEP Ingestion Analysis/custom_gigs_report_2026-06-09_01-34-14.xlsx"

wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("Sheet names:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("Sheet " + sheet_name + " is empty")
        continue

    headers = rows[0]
    data_rows = rows[1:]
    total_rows = len(data_rows)

    print("=== Sheet: " + sheet_name + " ===")
    print("Headers: " + str(headers))
    print("Total data rows: " + str(total_rows))

    for i, h in enumerate(headers):
        if h and 'function' in str(h).lower():
            print("Job function column found: " + repr(h) + " (index " + str(i) + ")")
            values = [row[i] for row in data_rows]
            blank_count = sum(1 for v in values if v is None or str(v).strip() == '')
            non_blank = [v for v in values if v is not None and len(str(v).strip()) > 0]
            counter = Counter(non_blank)
            print("Blank/null count: " + str(blank_count))
            print("Unique value counts:")
            for val, cnt in sorted(counter.items(), key=lambda x: -x[1]):
                print("  " + repr(val) + ": " + str(cnt))

    relevant_keywords = ['category', 'type', 'sector', 'domain', 'occupation', 'role', 'stream', 'family', 'function', 'field', 'area', 'job', 'cluster']
    print("All columns:")
    for i, h in enumerate(headers):
        if h:
            flag = " <-- RELEVANT" if any(kw in str(h).lower() for kw in relevant_keywords) else ""
            print("  [" + str(i) + "] " + str(h) + flag)
